import os
import io
import csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .db import get_db

def generate_excel_export(project_id: str) -> io.BytesIO:
    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT * FROM projects WHERE id = ?", (project_id,))
    proj = cur.fetchone()
    if not proj:
        conn.close()
        raise ValueError(f"Project {project_id} not found")

    cur.execute("""
    SELECT s.name as symbol_name, s.category, s.default_unit_rate,
           COUNT(d.id) as total_detected,
           SUM(CASE WHEN d.status IN ('auto', 'kept', 'manual') THEN 1 ELSE 0 END) as verified_qty
    FROM symbols s
    LEFT JOIN detections d ON s.id = d.symbol_id AND d.project_id = ?
    GROUP BY s.id, s.name, s.category, s.default_unit_rate
    """, (project_id,))
    summary_rows = cur.fetchall()

    cur.execute("""
    SELECT d.id, sh.name as sheet_name, d.grid_ref, s.name as symbol_name,
           d.confidence, d.status, d.bbox_x0, d.bbox_y0, d.bbox_x1, d.bbox_y1
    FROM detections d
    JOIN sheets sh ON d.sheet_id = sh.id
    JOIN symbols s ON d.symbol_id = s.id
    WHERE d.project_id = ?
    ORDER BY d.sheet_id, d.grid_ref
    """, (project_id,))
    detection_rows = cur.fetchall()

    cur.execute("""
    SELECT created_at, username, action, detection_id, details
    FROM audit_logs
    WHERE project_id = ?
    ORDER BY created_at DESC
    """, (project_id,))
    audit_rows = cur.fetchall()
    conn.close()

    wb = Workbook()

    # Style definitions
    navy_fill = PatternFill(start_color="0B3D91", end_color="0B3D91", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0B3D91")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # --- TAB 1: Takeoff Summary ---
    ws1 = wb.active
    ws1.title = "Takeoff Summary"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "ELECTROSCAN AI — ELECTRICAL TAKEOFF REPORT"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Project: {proj['name']} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Estimator: Parth Kala"
    ws1["A2"].font = subtitle_font

    headers1 = ["Symbol Name", "Category", "Verified Qty", "Unit Rate (₹)", "Total Amount (₹)"]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx > 2 else "left")

    row_num = 5
    for r in summary_rows:
        qty = r["verified_qty"] or 0
        rate = r["default_unit_rate"] or 0.0
        amount = qty * rate

        ws1.cell(row=row_num, column=1, value=r["symbol_name"]).font = regular_font
        ws1.cell(row=row_num, column=2, value=r["category"]).font = regular_font
        c_qty = ws1.cell(row=row_num, column=3, value=qty)
        c_qty.font = regular_font
        c_qty.alignment = Alignment(horizontal="right")
        
        c_rate = ws1.cell(row=row_num, column=4, value=rate)
        c_rate.font = regular_font
        c_rate.number_format = '"₹"#,##0.00'
        c_rate.alignment = Alignment(horizontal="right")

        c_amt = ws1.cell(row=row_num, column=5, value=amount)
        c_amt.font = bold_font
        c_amt.number_format = '"₹"#,##0.00'
        c_amt.alignment = Alignment(horizontal="right")

        for c in range(1, 6):
            ws1.cell(row=row_num, column=c).border = thin_border
        row_num += 1

    # Total row
    ws1.cell(row=row_num, column=1, value="TOTAL ESTIMATE").font = bold_font
    ws1.cell(row=row_num, column=3, value=f"=SUM(C5:C{row_num-1})").font = bold_font
    ws1.cell(row=row_num, column=5, value=f"=SUM(E5:E{row_num-1})").font = bold_font
    ws1.cell(row=row_num, column=5).number_format = '"₹"#,##0.00'
    for c in range(1, 6):
        ws1.cell(row=row_num, column=c).border = thin_border

    # --- TAB 2: Itemized Detections ---
    ws2 = wb.create_sheet(title="Itemized Detections")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = ["ID", "Sheet", "Grid Ref", "Symbol", "Confidence", "Status", "BBox (x0, y0, x1, y1)"]
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for idx, d in enumerate(detection_rows, 2):
        ws2.cell(row=idx, column=1, value=d["id"]).font = regular_font
        ws2.cell(row=idx, column=2, value=d["sheet_name"]).font = regular_font
        ws2.cell(row=idx, column=3, value=d["grid_ref"]).font = bold_font
        ws2.cell(row=idx, column=4, value=d["symbol_name"]).font = regular_font
        c_conf = ws2.cell(row=idx, column=5, value=d["confidence"])
        c_conf.font = regular_font
        c_conf.number_format = '0.0%'
        ws2.cell(row=idx, column=6, value=d["status"].upper()).font = bold_font
        ws2.cell(row=idx, column=7, value=f"[{d['bbox_x0']}, {d['bbox_y0']}, {d['bbox_x1']}, {d['bbox_y1']}]").font = regular_font
        for c in range(1, 8):
            ws2.cell(row=idx, column=c).border = thin_border

    # --- TAB 3: Audit Trail ---
    ws3 = wb.create_sheet(title="Audit Trail")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = ["Timestamp", "User", "Action", "Detection ID", "Details"]
    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for idx, a in enumerate(audit_rows, 2):
        ws3.cell(row=idx, column=1, value=str(a["created_at"])).font = regular_font
        ws3.cell(row=idx, column=2, value=a["username"]).font = regular_font
        ws3.cell(row=idx, column=3, value=a["action"]).font = bold_font
        ws3.cell(row=idx, column=4, value=a["detection_id"] or "-").font = regular_font
        ws3.cell(row=idx, column=5, value=a["details"]).font = regular_font
        for c in range(1, 6):
            ws3.cell(row=idx, column=c).border = thin_border

    # Auto-adjust column widths across all sheets
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_csv_export(project_id: str) -> str:
    conn = get_db()
    cur = conn.cursor()
    cur.execute("""
    SELECT d.id, sh.name as sheet, d.grid_ref, s.name as symbol, d.confidence, d.status,
           d.bbox_x0, d.bbox_y0, d.bbox_x1, d.bbox_y1
    FROM detections d
    JOIN sheets sh ON d.sheet_id = sh.id
    JOIN symbols s ON d.symbol_id = s.id
    WHERE d.project_id = ?
    ORDER BY d.sheet_id, d.grid_ref
    """, (project_id,))
    rows = cur.fetchall()
    conn.close()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Detection ID", "Sheet", "Grid Ref", "Symbol", "Confidence", "Status", "BBox X0", "BBox Y0", "BBox X1", "BBox Y1"])
    for r in rows:
        writer.writerow([r["id"], r["sheet"], r["grid_ref"], r["symbol"], f"{r['confidence']:.2%}", r["status"], r["bbox_x0"], r["bbox_y0"], r["bbox_x1"], r["bbox_y1"]])
    return output.getvalue()

def generate_pdf_report(project_id: str) -> io.BytesIO:
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM projects WHERE id = ?", (project_id,))
    proj = cur.fetchone()

    cur.execute("""
    SELECT s.name, s.category, s.default_unit_rate,
           COUNT(d.id) as total_detected,
           SUM(CASE WHEN d.status IN ('auto', 'kept', 'manual') THEN 1 ELSE 0 END) as verified_qty
    FROM symbols s
    LEFT JOIN detections d ON s.id = d.symbol_id AND d.project_id = ?
    GROUP BY s.id, s.name, s.category, s.default_unit_rate
    """, (project_id,))
    summary = cur.fetchall()
    conn.close()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        textColor=colors.HexColor('#0B3D91'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15
    )

    elements.append(Paragraph("ElectroScan AI — Official Takeoff Report", title_style))
    elements.append(Paragraph(f"Project: <b>{proj['name']}</b> &nbsp;|&nbsp; Date: {datetime.now().strftime('%B %d, %Y')} &nbsp;|&nbsp; Lead Estimator: Parth Kala", subtitle_style))
    elements.append(Spacer(1, 10))

    # Table
    table_data = [["Symbol", "Category", "Verified Count", "Unit Rate (₹)", "Total Amount (₹)"]]
    grand_total = 0.0
    for r in summary:
        qty = r["verified_qty"] or 0
        rate = r["default_unit_rate"] or 0.0
        amt = qty * rate
        grand_total += amt
        table_data.append([
            r["name"],
            r["category"],
            str(qty),
            f"₹{rate:.2f}",
            f"₹{amt:,.2f}"
        ])
    table_data.append(["TOTAL ESTIMATED COST", "", "", "", f"₹{grand_total:,.2f}"])

    t = Table(table_data, colWidths=[180, 100, 80, 80, 100])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0B3D91')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ALIGN', (2, 0), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#F1F5F9')),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 25))

    footer_text = Paragraph(
        "<i>Generated by ElectroScan AI. All quantities verified with click-to-source traceable audit trail.</i>",
        styles['Normal']
    )
    elements.append(footer_text)

    doc.build(elements)
    buf.seek(0)
    return buf
